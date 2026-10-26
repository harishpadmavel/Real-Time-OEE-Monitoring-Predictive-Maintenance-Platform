import docx, os

filename = 'Real-Time_OEE_Monitoring_and_Predictive_Maintenance_Platform_Report.docx'
doc = docx.Document(filename)

print('File size:', os.path.getsize(filename), 'bytes')
print('Sections:', len(doc.sections))
print('Paragraphs:', len(doc.paragraphs))
print('Tables:', len(doc.tables))

page_breaks = 0
for p in doc.paragraphs:
    for r in p.runs:
        if 'w:br' in r._r.xml and 'type="page"' in r._r.xml:
            page_breaks += 1

print('Explicit page breaks found:', page_breaks)

fonts = set()
font_sizes = set()
for p in doc.paragraphs:
    for r in p.runs:
        if r.font.name:
            fonts.add(r.font.name)
        if r.font.size:
            font_sizes.add(r.font.size.pt)

for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.font.name:
                        fonts.add(r.font.name)
                    if r.font.size:
                        font_sizes.add(r.font.size.pt)

print('Fonts used:', fonts)
print('Font sizes used:', sorted(list(font_sizes)))

# Check images embedded
num_images = 0
for p in doc.paragraphs:
    if 'w:drawing' in p._p.xml:
        num_images += 1
print('Embedded images count:', num_images)

print('\n--- Page-by-Page Headings Check ---')
page_num = 1
for p in doc.paragraphs:
    if p.text.strip() and p.runs and p.runs[0].bold and p.runs[0].font.size and p.runs[0].font.size.pt == 14:
        print(f"Page ~{page_num} heading: {p.text.strip()[:60]}")
    for r in p.runs:
        if 'w:br' in r._r.xml and 'type="page"' in r._r.xml:
            page_num += 1
            if page_num <= 12:
                print(f"---> Page Break to Page {page_num}")
